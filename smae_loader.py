"""
Importa los alimentos del archivo SMAE-5.xlsx a la base de datos SQLite.
Cada hoja del libro corresponde a un grupo de alimentos (Verduras, Frutas, etc.).
"""
import os

import openpyxl

import database

XLSX_NAME = "SMAE-5.xlsx"

COLUMNAS = [
    ("grupo", "grupo"),
    ("alimento", "alimento"),
    ("cantidad_sugerida", "cantidad"),
    ("unidad", "unidad"),
    ("peso_bruto_g", "peso_bruto"),
    ("peso_neto_g", "peso_neto"),
    ("energia_kcal", "energia"),
    ("proteina_g", "proteina"),
    ("lipidos_g", "lipidos"),
    ("hc_g", "hc"),
    ("fibra_g", "fibra"),
    ("vitamina_a_ug", "vitamina_a"),
    ("vitamina_c_mg", "vitamina_c"),
    ("acido_folico_ug", "acido_folico"),
    ("hierro_mg", "hierro"),
    ("potasio_mg", "potasio"),
    ("calcio_mg", "calcio"),
    ("sodio_mg", "sodio"),
    ("azucar_g", "azucar"),
    ("indice_glucemico", "ig"),
    ("carga_glucemica", "cg"),
]

# Palabras clave (en minúsculas) para localizar la columna de cada nutriente.
BUSQUEDAS = {
    "alimento": ["alimento"],
    "cantidad": ["cantidad sugerida", "cantidad"],
    "unidad": ["unidad"],
    "peso_bruto": ["peso bruto", "peso redondo", "peso redondeado"],
    "peso_neto": ["peso neto"],
    "energia": ["energia", "energía", "enegia"],
    "proteina": ["proteina", "proteína", "proteina"],
    "lipidos": ["lipido", "lípido"],
    "hc": ["hidrato", "carbono", " hc"],
    "fibra": ["fibra"],
    "vitamina_a": ["vitamina a", "vit "],
    "vitamina_c": ["ascorbico", "ascórbico"],
    "acido_folico": ["folico", "fólico"],
    "hierro": ["hierro"],
    "potasio": ["potasio"],
    "calcio": ["calcio"],
    "sodio": ["sodio"],
    "azucar": ["azucar", "azúcar"],
    "ig": ["indice glic", "índice glic", "indice gluc", "glucemico"],
    "cg": ["carga glic", "carga gluc"],
}

NUTRIENTES_NUMERICOS = [clave for _, clave in COLUMNAS if clave != "grupo"]


def _a_float(valor):
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(",", ".")
    if texto.upper() in ("", "ND", "N/D", "NA", "N/A", "-"):
        return 0.0
    try:
        return float(texto)
    except ValueError:
        return 0.0


def _mapear_columnas(encabezados):
    mapa = {}
    normalizados = [
        (idx, str(val).strip().lower()) for idx, val in enumerate(encabezados)
    ]
    for campo, claves in BUSQUEDAS.items():
        for idx, texto in normalizados:
            if idx in mapa.values():
                continue
            if any(clave in texto for clave in claves):
                mapa[campo] = idx
                break
    return mapa


def _leer_alimentos():
    ruta = os.path.join(os.path.dirname(os.path.abspath(__file__)), XLSX_NAME)
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"No se encuentra {ruta}")

    wb = openpyxl.load_workbook(ruta, data_only=True)
    filas = []
    for nombre_hoja in wb.sheetnames:
        ws = wb[nombre_hoja]
        encabezados = None
        fila_datos = None
        for idx, fila in enumerate(
            ws.iter_rows(min_row=1, max_row=min(10, ws.max_row), values_only=True),
            start=1,
        ):
            textos = [str(c).strip().lower() if c is not None else "" for c in fila]
            if any("alimento" in t for t in textos):
                encabezados = list(fila)
                fila_datos = idx + 1
                break
        if encabezados is None:
            continue

        mapa = _mapear_columnas(encabezados)
        if "alimento" not in mapa:
            continue

        for fila in ws.iter_rows(
            min_row=fila_datos, max_row=ws.max_row, values_only=True
        ):
            def valor(campo):
                idx = mapa.get(campo)
                if idx is None or idx >= len(fila):
                    return None
                return fila[idx]

            alimento = valor("alimento")
            alimento = str(alimento).strip() if alimento is not None else ""
            if not alimento:
                continue

            filas.append(
                {
                    "grupo": nombre_hoja,
                    "alimento": alimento,
                    "cantidad": _a_float(valor("cantidad")),
                    "unidad": (
                        str(valor("unidad")).strip()
                        if valor("unidad") is not None
                        else ""
                    ),
                    "peso_bruto": _a_float(valor("peso_bruto")),
                    "peso_neto": _a_float(valor("peso_neto")),
                    "energia": _a_float(valor("energia")),
                    "proteina": _a_float(valor("proteina")),
                    "lipidos": _a_float(valor("lipidos")),
                    "hc": _a_float(valor("hc")),
                    "fibra": _a_float(valor("fibra")),
                    "vitamina_a": _a_float(valor("vitamina_a")),
                    "vitamina_c": _a_float(valor("vitamina_c")),
                    "acido_folico": _a_float(valor("acido_folico")),
                    "hierro": _a_float(valor("hierro")),
                    "potasio": _a_float(valor("potasio")),
                    "calcio": _a_float(valor("calcio")),
                    "sodio": _a_float(valor("sodio")),
                    "azucar": _a_float(valor("azucar")),
                    "ig": _a_float(valor("ig")),
                    "cg": _a_float(valor("cg")),
                }
            )
    wb.close()
    return filas


def importar_smae():
    filas = _leer_alimentos()
    database.recrear_smae_alimentos()
    campos_db = [columna for columna, _ in COLUMNAS]
    claves = [clave for _, clave in COLUMNAS]
    with database.get_connection() as conn:
        columnas = ", ".join(campos_db)
        marcadores = ", ".join(":" + c for c in claves)
        conn.executemany(
            f"INSERT INTO smae_alimentos ({columnas}) VALUES ({marcadores})",
            [
                {clave: fila.get(clave, 0) for clave in claves}
                for fila in filas
            ],
        )
    return len(filas)


if __name__ == "__main__":
    total = importar_smae()
    print(f"Importados {total} alimentos a smae_alimentos desde {XLSX_NAME}")
